"""Excel exports: bold coloured header, one sheet per kind of data, timestamped file name."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from conftest import prediction_item


def test_history_export(client, user, features):
    client.post("/db/predictions", json={"items": [prediction_item(features)]}, headers=user["headers"])
    res = client.get("/db/export/predictions.xlsx", headers=user["headers"])
    assert res.status_code == 200
    assert 'filename="lich_su_du_doan_' in res.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Lịch sử dự đoán", "Đặc trưng đầu vào", "Bộ lọc"]
    ws = wb["Lịch sử dự đoán"]
    assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("6C5CE7")
    assert ws.max_row == 2 and ws["E2"].value == "Ác tính"
    assert wb["Đặc trưng đầu vào"].max_column == 31


def test_model_comparison_export(client, user):
    res = client.get("/db/export/model-runs.xlsx", headers=user["headers"])
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["So sánh mô hình", "Lịch sử train"]
    ws = wb["So sánh mô hình"]
    assert ws["B1"].value == "Sensitivity" and ws.max_row >= 6


def test_exports_need_login(client):
    assert client.get("/db/export/predictions.xlsx").status_code == 401
    assert client.get("/db/export/model-runs.xlsx").status_code == 401
