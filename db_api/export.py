"""Excel exports (openpyxl): prediction history and the model comparison table.

Every workbook has a bold, coloured header row, auto-sized columns and one sheet
per kind of data; file names carry a timestamp.
"""

from __future__ import annotations

import io
import json
from datetime import UTC, date, datetime

from fastapi import APIRouter, Depends, HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import security
from db_api.db import UTC_FORMAT, get_conn, local_tz
from db_api.history import history_filter

router = APIRouter(prefix="/export", tags=["Xuất Excel"])

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="6C5CE7")
BEST_FILL = PatternFill("solid", fgColor="EFEBFF")
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LABEL_VI = {"malignant": "Ác tính", "benign": "Lành tính", None: ""}
WARNING = "Chỉ phục vụ mục đích giáo dục, không thay thế chẩn đoán y khoa."


def to_local(utc_iso: str | None) -> str:
    """'2026-09-24T02:00:00Z' -> '2026-09-24 09:00:00' in APP_TZ_OFFSET_HOURS."""
    if not utc_iso:
        return ""
    dt = datetime.strptime(utc_iso, UTC_FORMAT).replace(tzinfo=UTC)
    return dt.astimezone(local_tz()).strftime("%Y-%m-%d %H:%M:%S")


def filename(prefix: str) -> str:
    return f"{prefix}_{datetime.now(local_tz()):%Y%m%d_%H%M%S}.xlsx"


def _write_sheet(ws, headers: list[str], rows: list[list], highlight_rows: set[int] | None = None) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for idx in highlight_rows or ():
        for cell in ws[idx + 2]:
            cell.fill = BEST_FILL
    ws.freeze_panes = "A2"
    # Auto-fit: widest rendered value per column, capped so long columns stay readable.
    for col_idx, column in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in column)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 8), 60)


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def predictions_workbook(rows: list, username: str, filters: dict) -> bytes:
    """Sheet 1: one row per prediction. Sheet 2: the 30 input values. Sheet 3: filters used."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Lịch sử dự đoán"
    headers = ["ID", "Thời gian", "Mô hình", "Mẫu", "Dự đoán", "P(ác tính)", "Ngưỡng", "Nhãn thật",
               "Đúng / sai", "Thời gian chạy (ms)"]
    data = []
    for r in rows:
        correct = "" if r["actual_label"] is None else ("Đúng" if r["actual_label"] == r["predicted_label"] else "Sai")
        data.append([
            r["id"], to_local(r["created_at"]), r["model"], r["sample_id"] or "",
            LABEL_VI[r["predicted_label"]], r["probability_malignant"], r["threshold"],
            LABEL_VI[r["actual_label"]], correct,
            round(r["runtime_ms"], 4) if r["runtime_ms"] is not None else None,
        ])
    _write_sheet(ws, headers, data)

    features = wb.create_sheet("Đặc trưng đầu vào")
    inputs = [json.loads(r["input_json"]) for r in rows]
    names = list(inputs[0]) if inputs else []
    _write_sheet(features, ["ID", *names], [[r["id"], *[i.get(n) for n in names]] for r, i in zip(rows, inputs,
                                                                                                   strict=True)])

    info = wb.create_sheet("Bộ lọc")
    _write_sheet(info, ["Thông tin", "Giá trị"], [
        ["Người dùng", username],
        ["Xuất lúc", datetime.now(local_tz()).strftime("%Y-%m-%d %H:%M:%S")],
        ["Lọc mô hình", filters.get("model") or "Tất cả"],
        ["Từ ngày", filters.get("date_from") or "—"],
        ["Đến ngày", filters.get("date_to") or "—"],
        ["Số dòng", len(rows)],
        ["Cảnh báo", WARNING],
    ])
    return _to_bytes(wb)


MODEL_HEADERS = ["Mô hình", "Sensitivity", "Specificity", "Precision ác tính", "F1 ác tính", "ROC-AUC", "PR-AUC",
                 "Brier", "Ngưỡng", "FN", "FP", "CV recall_macro (TB)", "CV (độ lệch)", "Thời gian train (s)",
                 "Thời gian predict (ms/mẫu)", "Tham số tốt nhất", "Tốt nhất"]


def _model_row(m) -> list:
    return [
        m["label"] or m["model"], m["sensitivity"], m["specificity"], m["precision_malignant"], m["f1_malignant"],
        m["roc_auc"], m["pr_auc"], m["brier"], m["threshold"], m["false_negatives"], m["false_positives"],
        m["cv_mean"], m["cv_std"], m["train_seconds"], m["predict_ms_per_sample"], m["best_params_json"],
        "✔" if m["is_best"] else "",
    ]


def model_runs_workbook(run, models: list, all_runs: list) -> bytes:
    """Sheet 1: the chosen run's comparison table. Sheet 2: every run, one row per model."""
    wb = Workbook()
    ws = wb.active
    ws.title = "So sánh mô hình"
    best_idx = {i for i, m in enumerate(models) if m["is_best"]}
    _write_sheet(ws, MODEL_HEADERS, [_model_row(m) for m in models], best_idx)
    ws.append([])
    ws.append([f"Lần train #{run['id']} lúc {to_local(run['created_at'])} — {run['train_size']} mẫu train / "
               f"{run['test_size']} mẫu test; ngưỡng mỗi mô hình chọn cho sensitivity out-of-fold ≥ "
               f"{run['target_sensitivity']}; mô hình tốt nhất chọn theo CV recall_macro."])
    ws.append([WARNING])

    history = wb.create_sheet("Lịch sử train")
    rows = [[m["run_id"], to_local(m["run_created_at"]), *_model_row(m)] for m in all_runs]
    _write_sheet(history, ["Lần train", "Thời gian", *MODEL_HEADERS], rows)
    return _to_bytes(wb)


# ----------------------------------------------------------------- routes


def _download(content: bytes, name: str) -> Response:
    return Response(content, media_type=XLSX_MIME,
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/predictions.xlsx")
def export_predictions(
    model: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    user: dict = Depends(security.current_user),
    conn=Depends(get_conn),
):
    """The current user's history (same filters as GET /predictions) as an .xlsx file."""
    where, args = history_filter(user["id"], model, date_from, date_to)
    rows = conn.execute(f"SELECT * FROM predictions WHERE {where} ORDER BY created_at DESC, id DESC", args).fetchall()
    filters = {"model": model, "date_from": date_from and date_from.isoformat(),
               "date_to": date_to and date_to.isoformat()}
    return _download(predictions_workbook(rows, user["username"], filters), filename("lich_su_du_doan"))


@router.get("/model-runs.xlsx")
def export_model_runs(run_id: int | None = None, user: dict = Depends(security.current_user),
                      conn=Depends(get_conn)):
    """Comparison table of one run (latest by default) plus the history of every run."""
    if run_id is None:
        run = conn.execute("SELECT * FROM training_runs ORDER BY id DESC LIMIT 1").fetchone()
    else:
        run = conn.execute("SELECT * FROM training_runs WHERE id = ?", (run_id,)).fetchone()
    if run is None:
        raise HTTPException(404, "Chưa có lần train nào được lưu")
    models = conn.execute("SELECT * FROM model_runs WHERE run_id = ? ORDER BY id", (run["id"],)).fetchall()
    all_runs = conn.execute(
        """SELECT r.id AS run_id, r.created_at AS run_created_at, m.* FROM training_runs r
           JOIN model_runs m ON m.run_id = r.id ORDER BY r.id DESC, m.id""").fetchall()
    return _download(model_runs_workbook(run, models, all_runs), filename("so_sanh_mo_hinh"))
